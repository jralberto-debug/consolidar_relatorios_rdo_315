import os
import pandas as pd
import openpyxl

# Caminho da pasta com os arquivos RDO
pasta = r"C:\Users\Engevale\Documents\OneDrive\315.[TANQUES][ENGEVALE][REPAR][ONE]\315\RDO"
saida = os.path.join(pasta, "consolidado_relatorios.xlsx")

dados = []

# Verificar arquivos que começam com RDO
for arquivo in os.listdir(pasta):
    if not arquivo.startswith("RDO") or not arquivo.endswith(".xlsx"):
        continue

    caminho_arquivo = os.path.join(pasta, arquivo)
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)

    print(f"🔍 Lendo arquivo: {arquivo}")

    for aba_nome in wb.sheetnames:
        if "TQ" not in aba_nome.upper():
            continue

        print(f"   ➤ Processando aba: {aba_nome}")
        aba = wb[aba_nome]
        tanque = aba_nome

        # Pega dados fixos
        relatorio_num = aba.cell(row=1, column=13).value  # M1
        periodo = aba.cell(row=2, column=9).value         # I2

        if isinstance(periodo, str):
            periodo = periodo.strip()

        linha = 7
        linhas_vazias = 0
        dados_encontrados = 0

        while True:
            valores = [aba.cell(row=linha, column=col).value for col in range(1, 14)]
            if all(v is None or str(v).strip() == '' for v in valores[:3]):
                linhas_vazias += 1
                if linhas_vazias >= 5:
                    break
                linha += 1
                continue

            linhas_vazias = 0
            dados_encontrados += 1

            item, servico, _, met_dia = valores[0], valores[1], valores[2], valores[3]
            seg, ter, qua, qui, sex, sab, dom = valores[4:11]
            prev, real = valores[11], valores[12]

            dados.append([
                tanque, relatorio_num, periodo,
                item, servico, met_dia,
                seg, ter, qua, qui, sex, sab, dom,
                prev, real
            ])
            linha += 1

        if dados_encontrados == 0:
            print(f"   ⚠️ Nenhuma linha útil encontrada na aba {aba_nome}")

# Salvar planilha final
df = pd.DataFrame(dados, columns=[
    "TANQUE", "RELATÓRIO Nº", "PERÍODO",
    "ITEM", "SERVIÇO", "MÉT. DIA",
    "SEG", "TER", "QUA", "QUI", "SEX", "SÁB", "DOM",
    "ACUM. PREV", "ACUM. REAL"
])

df.to_excel(saida, index=False)
print(f"\n✅ Arquivo gerado com sucesso: {saida}")
